from pathlib import Path

from docx import Document

from app.ingestion.hashing import hash_bytes
from app.ingestion.loaders.base import PageContent, RawDocument


def load_docx(path: Path) -> RawDocument:
    raw_bytes = path.read_bytes()
    content_hash = hash_bytes(raw_bytes)

    doc = Document(path)  # python-docx needs a path or file-like object, not raw bytes
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    return RawDocument(
        source_type="docx",
        source_name=path.name,
        content_hash=content_hash,
        pages=[PageContent(page_number=1, text=text)],
    )